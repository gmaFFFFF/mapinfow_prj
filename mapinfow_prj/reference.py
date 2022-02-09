from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.table import Table

from .xlsx_table_with_index import ТаблицаСИндексом


class СкПараметрыXlsx:
    """
    Файл xlsx с параметрами систем координат, используемых MapInfo
    """

    @dataclass(frozen=True)
    class Эллипсоид:
        код: int
        имя: str
        proj: str

    @dataclass(frozen=True)
    class Датум:
        код: int
        имя: str
        proj: str
        ellps: str
        towgs84: str
        pm: str

    @dataclass(frozen=True)
    class ЕдИзм:
        код: int
        имя: str
        proj_code: str

    @dataclass(frozen=True)
    class Проекция:
        код: int
        имя: str
        число_парам: int
        proj_format: str

    путь: Path = Path(__file__).parent / 'cs_params.xlsx'

    def __init__(self, путь=None) -> None:
        super().__init__()

        if путь is not None:
            self.путь = путь

        self.книга: Workbook = load_workbook(self.путь, data_only=True)

    def _выгрузить_данные_из_умной_таблицы(self, название: str, индекс: str):
        табл: Table = None
        лист, табл = [(лист, лист.tables[табл])
                      for лист in self.книга.worksheets
                      for табл in лист.tables.keys() if название == табл][0]
        заголовки: list[str] = табл.column_names
        ячейки: tuple[tuple[Cell]] = лист[табл.ref][1:]
        данные = [[ячейка.value for ячейка in стр] for стр in ячейки]

        return ТаблицаСИндексом(заголовки, данные, индекс)

    def дай_эллипсоиды(self) -> list[Эллипсоид]:
        табл = self._выгрузить_данные_из_умной_таблицы("ellips", "code")
        return [СкПараметрыXlsx.Эллипсоид(стр["code"], стр["name"], стр["proj"]) for стр in табл]

    def дай_датумы(self) -> list[Датум]:
        табл = self._выгрузить_данные_из_умной_таблицы("datum", "code")

        # В xlsx файл учтена необходимость изменить знак параметров поворота у +towgs84=
        # вызов функции _изменить_знак_параметров_поворота_towgs84() не требуется

        return [СкПараметрыXlsx.Датум(стр["code"], стр["name"], стр["proj"],
                                      стр["ellps"], стр["towgs84"], стр["pm"]) for стр in табл]

    def дай_ед_изм(self) -> list[ЕдИзм]:
        табл = self._выгрузить_данные_из_умной_таблицы("unit", "code")
        return [СкПараметрыXlsx.ЕдИзм(стр["code"], стр["name"], стр["proj_code"]) for стр in табл]

    def дай_проекции(self) -> list[Проекция]:
        табл = self._выгрузить_данные_из_умной_таблицы("prj", "code")
        return [
            СкПараметрыXlsx.Проекция(стр["code"], стр["name"], стр["params_cnt"], стр["proj_format"]) for стр in табл
        ]

    @staticmethod
    def _изменить_знак_параметров_поворота_towgs84(towgs84: str):
        """
        В ГОСТ 32453-2017 и MapInfo Pro используется формула "Поворот координатной системы".
        А в ISO 19111 и в библиотеке Proj (ранее Proj.4) используется "Преобразование радиуса-вектора".
        Таким образом, чтобы использовать параметры датумов ГОСТ и MapInfo в Proj требуется изменить
        знаки у параметров поворота ωx, ωy, ωz.
        Источник: https://mapinfo.ru/articles/gost

        >>> СкПараметрыXlsx._изменить_знак_параметров_поворота_towgs84('+towgs84=23.92,-141.27,-80.9,0,-0.35,-0.82,-0.12')
        '+towgs84=23.92,-141.27,-80.9,0,0.35,0.82,-0.12'

        :param towgs84: строка вида '+towgs84=23.92,-141.27,-80.9,0,-0.35,-0.82,-0.12'

        :return: строка вида '+towgs84=23.92,-141.27,-80.9,0,0.35,0.82,-0.12', с измененным знаком параметров поворота
        """

        стр_парам = towgs84.removeprefix("+towgs84=")
        парам = [int(п.strip()) if п.strip().isdigit() else float(п.strip())
                 for п in стр_парам.split(",")]
        if len(парам) == 3:
            return towgs84
        парам_коррект = [str(-п) if i in (3, 4, 5) else str(п)
                         for i, п in enumerate(парам)]
        return f'+towgs84={",".join(парам_коррект)}'
